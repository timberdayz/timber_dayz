"""商品中心 SKU 采购价批量录入脚本.

读取 product-catalog skill 的 sku-reverse-lookup.xlsx,批量写入
PostgreSQL core.dim_erp_sku.default_purchase_cost 及其元数据。

详见 docs/superpowers/specs/2026-09-18-purchase-price-import-design.md
"""
from __future__ import annotations

import argparse
import asyncio
import csv
import logging
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import openpyxl


SHEET_NAME = "SKU反查表"
COL_SKU_CODE = 0  # 第 1 列 (0-indexed)
COL_PRICE = 8     # 第 9 列 (0-indexed)
EM_DASH = "—"
SOURCE_VALUE = "妙手导入"
CONFIDENCE_VALUE = "medium"
CURRENCY_VALUE = "CNY"


@dataclass(frozen=True)
class ParsedRow:
    sku_code: str
    price: float


def parse_lookup_xlsx(path: str) -> tuple[list[tuple[str, float]], list[str], list[str]]:
    """读取 xlsx 并按过滤规则分类.

    Returns:
        (valid, skipped, errors)
        - valid:   [(sku_code, price), ...]   # 价格 > 0 的数字
        - skipped: [sku_code, ...]             # 被跳过的 SKU code(— / 0 / 空)
        - errors:  [message, ...]              # 解析错误信息(字符串型非 — 等)
    """
    valid: list[tuple[str, float]] = []
    skipped: list[str] = []
    errors: list[str] = []

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        errors.append(f"xlsx 找不到 sheet '{SHEET_NAME}',实际 sheets: {wb.sheetnames}")
        return valid, skipped, errors
    ws = wb[SHEET_NAME]

    rows = ws.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        errors.append("xlsx 没有表头行")
        return valid, skipped, errors

    if len(header) < 9 or header[COL_SKU_CODE] != "SKU code" or header[COL_PRICE] != "采购单价(元)":
        errors.append(
            f"xlsx 表头不符,期望第1列='SKU code' 第9列='采购单价(元)',实际: {header[:9]}"
        )
        return valid, skipped, errors

    for row in rows:
        if row is None:
            continue
        sku_code = row[COL_SKU_CODE]
        price = row[COL_PRICE]
        if sku_code is None or str(sku_code).strip() == "":
            continue
        sku_code_str = str(sku_code).strip()
        if price is None or (isinstance(price, str) and price.strip() == ""):
            skipped.append(sku_code_str)
            continue
        if isinstance(price, str):
            if price.strip() == EM_DASH:
                skipped.append(sku_code_str)
                continue
            errors.append(f"{sku_code_str}: 价格是字符串 '{price=}' 非数字,且非 — ")
            continue
        if isinstance(price, (int, float)):
            if price <= 0:
                skipped.append(sku_code_str)
                continue
            valid.append((sku_code_str, float(price)))
            continue
        errors.append(f"{sku_code_str}: 价格类型未知 {type(price).__name__}")

    return valid, skipped, errors